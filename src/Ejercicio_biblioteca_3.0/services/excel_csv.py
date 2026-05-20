import csv
from openpyxl import load_workbook
import os

def importar_recursos_excel(ruta :str) -> list:
    if not os.path.exists(ruta):
        raise FileNotFoundError(f"El archivo '{ruta}' no existe.")
    
    workbook = load_workbook(ruta)
    hoja = workbook.active
    nuevos_recursos = []

    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if not fila[0]:
            continue

        datos = {
            "id": str(fila[0]),
            "tipo": str(fila[1]).strip().lower(),  # 'libro', 'revista', 'tesis', 'digital'
            "titulo": str(fila[2]),
            "autor": str(fila[3]),
            "anyo": int(fila[4]),
            "extra_1": fila[5],  # Páginas, Edición, Universidad o Formato
            "extra_2": fila[6]   # Tamaño MB (solo para digital)
        }
        nuevos_recursos.append(datos)

    return nuevos_recursos

def exportar_recursos_csv(ruta : str, lista_recursos : list):
    os.makedirs(os.path.dirname(ruta), exist_ok=True)

    with open(ruta, "w", encoding="utf-8", newline="") as file:
        escritura = csv.writer(file)

        escritura.writerow(["ID", "Tipo", "Título", "Autor", "Año", "Prestado"])

        for recurso in lista_recursos:
            escritura.writerow([
                recurso.id,
                recurso.__class__.__name__,
                recurso.titulo,
                recurso.autor,
                recurso.anyo,
                "Sí" if recurso.prestado else "No"
            ])
    
def exportar_historial_usuario_csv(usuario_nombre: str, registros_historial: list):
    nombre_archivo = f"data/historial_{usuario_nombre.lower().replace(' ', '_')}.csv"

    with open(nombre_archivo, "w", encoding="utf-8", newline="") as file:
        escritura = csv.writer(file)
        escritura.writerow(["Fecha", "Recurso ID", "Título Recurso", "Acción"])

        for registro in registros_historial:
            if registro["usuario_nombre"].lower() == usuario_nombre.lower():
                escritura.writerow([
                    registro["fecha"],
                    registro["recurso_id"],
                    registro["titulo"],
                    registro["accion"]
                ])