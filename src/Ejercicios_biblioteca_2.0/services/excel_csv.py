import csv
from openpyxl import load_workbook
from models.libro import Libro

def importar_libros_excel(biblioteca, ruta):
    workbook = load_workbook(ruta)
    hoja = workbook.active

    for fila in hoja.iter_rows(min_row=2, values_only=True):
        titulo = fila[0]
        autor = fila[1]
        prestado = fila[2]

        libro = Libro(titulo,autor)
        libro.prestado = prestado
        biblioteca.agregar_libro(libro)

def exportar_libros_csv(lista_libros):
    with open("data/libros.csv", "w", encoding="utf-8") as file:
        escritura = csv.writer(file)

        escritura.writerow([
            "titulo",
            "autor",
            "prestado"
        ])
        for libro in lista_libros:
            escritura.writerow([
                libro.titulo,
                libro.autor,
                libro.prestado
            ])